"""Build the SIH26119 idea deck (6 slides) on the official SIH 2026 template, reusing idea-1's styling."""
import copy
import math
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

HERE = Path(__file__).parent
ROOT = HERE.parent
BLUE, INK, MUTED, CARD = "1F4E9A", "202020", "555555", "EEF3FA"
GREEN, GREEN_BG, AMBER_BG = "1E7B4F", "E6F4EC", "FFF4E0"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# results/final.json (bench_final.py): per size, HiGHS best of simplex/IPM, SAMADHAN best of its two
# GPU methods at 1e-4 ("fast") and 1e-6 ("acc") relative KKT.
bench = []
for r in json.loads((ROOT / "results" / "final.json").read_text()):
    hb = r["highs"]["best"]
    bench.append(dict(size=r["size"], n=r["n"], m=r["m"], nnz=r["nnz"],
                      highs=dict(hb, method=hb.get("method", "")),
                      fast=r["samadhan"]["0.0001"]["best"], acc=r["samadhan"]["1e-06"]["best"]))
byk = {r["size"]: r for r in bench}
big = byk[max(byk, key=lambda k: byk[k]["n"])]
solved = lambda r: r["highs"]["status"] == "Optimal"
speed = lambda r, t="acc": r["highs"]["time"] / r[t]["time"]
X = "×"


def spd_txt(r, t="acc"):
    s = speed(r, t)
    txt = f"{s:.1f}{X}" if s < 10 else f"{s:.0f}{X}"
    return txt if solved(r) else ">" + txt


speed_big = speed(big)
checked = [r for r in bench if solved(r)]
worst_acc = max(r["acc"]["obj_err"] for r in bench)
worst_fast = max(r["fast"]["obj_err"] for r in bench)

milp = json.loads((ROOT / "results" / "milp.json").read_text())
milp_ok = sum(1 for r in milp if r["status"] == "optimal" and r["rel_diff"] < 1e-9)

REPO = "https://github.com/NAMAN-THAKUR-1944/SIH2026-SIH26119-SAMADHAN"
netlib = None
_nl = ROOT / "results" / "netlib_cpu_tol0.0001.json"
if _nl.exists():
    netlib = json.loads(_nl.read_text())

prs = Presentation(str(HERE / "idea1.pptx"))
CHROME = ("Rectangle", "Title 1", "Slide Number", "Footer", "Oval")


def strip(slide, keep_names=()):
    for sh in list(slide.shapes):
        nm = sh.name
        chrome = nm.startswith(CHROME) or (sh.shape_type == 13 and sh.left > Inches(10.5) and sh.top < Inches(0.2))
        if not chrome and nm not in keep_names:
            sh._element.getparent().remove(sh._element)


def run_xml(text, size, bold=False, color=INK, italic=False):
    r = etree.SubElement(etree.Element(qn("a:dummy")), qn("a:r"))
    rpr = etree.SubElement(r, qn("a:rPr"), sz=str(int(size * 100)), b="1" if bold else "0", lang="en-US")
    if italic:
        rpr.set("i", "1")
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr"), val=color)
    etree.SubElement(rpr, qn("a:latin"), typeface="Arial")
    t = etree.SubElement(r, qn("a:t"))
    t.text = text
    return r


def para(kind, parts, size=None, before=None, align=None, color=None):
    """kind: 'h1' (big diamond header), 'h2' (diamond header), 'b' (bullet), 'p' (plain)."""
    p = etree.Element(qn("a:p"))
    ppr = etree.SubElement(p, qn("a:pPr"))
    if align:
        ppr.set("algn", align)
    if kind in ("h1", "h2"):
        ppr.set("marL", "228600"); ppr.set("indent", "-228600")
        if before is None:
            before = 0 if kind == "h1" else 700
    elif kind == "b":
        ppr.set("marL", "457200"); ppr.set("indent", "-228600")
        before = 250 if before is None else before
    else:
        ppr.set("marL", "0"); ppr.set("indent", "0")
    if before:
        sb = etree.SubElement(ppr, qn("a:spcBef")); etree.SubElement(sb, qn("a:spcPts"), val=str(before))
    if kind in ("h1", "h2"):
        etree.SubElement(ppr, qn("a:buFont"), typeface="Wingdings"); etree.SubElement(ppr, qn("a:buChar"), char="v")
        size = size or (1700 if kind == "h1" else 1500) / 100
        for txt in ([parts] if isinstance(parts, str) else parts):
            p.append(run_xml(txt, size, True, color or BLUE))
    else:
        if kind == "b":
            etree.SubElement(ppr, qn("a:buFont"), typeface="Arial"); etree.SubElement(ppr, qn("a:buChar"), char="•")
        else:
            etree.SubElement(ppr, qn("a:buNone"))
        size = size or 13
        if isinstance(parts, str):
            parts = [("", parts)]
        for bold_txt, txt in parts:
            if bold_txt:
                p.append(run_xml(bold_txt, size, True, color or INK))
            if txt:
                p.append(run_xml(txt, size, False, color or INK))
    return p


def textbox(slide, x, y, w, h, paras, anchor=None, autofit=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    body = tf._txBody
    for p in body.findall(qn("a:p")):
        body.remove(p)
    for p in paras:
        body.append(p)
    bp = body.find(qn("a:bodyPr"))
    for k in ("lIns", "rIns", "tIns", "bIns"):
        bp.set(k, "45720")
    if anchor:
        tf.vertical_anchor = anchor
    return tb


def card(slide, x, y, w, h, fill=CARD, line=None, paras=(), anchor=MSO_ANCHOR.TOP, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.07
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        s.line.color.rgb = RGBColor.from_string(line); s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    body = tf._txBody
    bp = body.find(qn("a:bodyPr"))
    for k, v in (("lIns", "100584"), ("rIns", "100584"), ("tIns", "64008"), ("bIns", "64008")):
        bp.set(k, v)
    for p in body.findall(qn("a:p")):
        body.remove(p)
    for p in paras:
        body.append(p)
    if not paras:  # a txBody must hold at least one paragraph
        etree.SubElement(body, qn("a:p"))
    return s


def stat(slide, x, y, w, h, big_txt, small_txt, size=24):
    card(slide, x, y, w, h, paras=[
        para("p", [(big_txt, "")], size=size, align="ctr", color=BLUE),
        para("p", small_txt, size=10.5, align="ctr", before=200)], anchor=MSO_ANCHOR.MIDDLE)


def caption(slide, x, y, w, h, txt):
    textbox(slide, x, y, w, h, [para("p", txt, size=10, color=MUTED)])


def set_title(slide, text, size=None):
    for sh in slide.shapes:
        if sh.name == "Title 1":
            tf = sh.text_frame
            r0 = tf.paragraphs[0].runs[0]
            r0.text = text
            for r in tf.paragraphs[0].runs[1:]:
                r._r.getparent().remove(r._r)
            for p in tf.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            if size:
                r0.font.size = Pt(size)
            return sh


def log_axis(chart, lo=0.1):
    scaling = chart.value_axis._element.find(qn("c:scaling"))
    lb = etree.SubElement(scaling, qn("c:logBase")); lb.set("val", "10")
    scaling.insert(0, lb)
    mn = etree.SubElement(scaling, qn("c:min")); mn.set("val", str(lo))
    # crosses at the floor so bars grow upward from it
    ax = chart.category_axis._element
    cr = ax.find(qn("c:crosses"))
    if cr is not None:
        cz = etree.Element(qn("c:crossesAt")); cz.set("val", str(lo)); cr.addprevious(cz); ax.remove(cr)


def htime_txt(r):
    return fmt_s(r["highs"]["time"]) if solved(r) else f">{r['highs']['time']:.0f} s (limit)"


def pct(x):
    """Percent with two significant digits, never in exponent form (e.g. 0.00035%, 0.070%)."""
    p = x * 100
    if p <= 0:
        return "0%"
    d = max(2, -math.floor(math.log10(p)) + 1)
    return f"{p:.{d}f}%"


def fmt_s(t):
    return f"{t:.1f} s" if t < 100 else f"{t:.0f} s"


SIZES = [r["size"] for r in bench]
SERIES_COLORS = ("A6A6A6", BLUE, "6FA8DC")
label = lambda r: f"{r['n']/1e6:.2f}M" if r["n"] >= 1e6 else f"{r['n']/1e3:.0f}k"


def time_chart(slide, x, y, w, h, title):
    cd = CategoryChartData()
    cd.categories = [f"{label(r)} vars" for r in bench]
    cd.add_series("HiGHS, CPU (best method)", [round(r["highs"]["time"], 2) for r in bench])
    cd.add_series("SAMADHAN GPU, accurate (1e-6)", [round(r["acc"]["time"], 2) for r in bench])
    cd.add_series("SAMADHAN GPU, fast (1e-4)", [round(r["fast"]["time"], 2) for r in bench])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    tp = ch.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size = Pt(12); tp.runs[0].font.bold = True; tp.runs[0].font.color.rgb = RGBColor.from_string(INK)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10)
    for s, col in zip(ch.series, SERIES_COLORS):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = RGBColor.from_string(col)
    pl = ch.plots[0]
    pl.gap_width = 50
    pl.overlap = 0
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.font.size = Pt(8); dl.number_format = '[<1]0.0" s";0" s"'; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor.from_string("E0E0E0")
    va.tick_labels.font.size = Pt(9)
    va.tick_labels.number_format = '[<1]0.0" s";0" s"'; va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    log_axis(ch)
    ch.category_axis.tick_labels.font.size = Pt(10)
    mark_timeouts(ch)
    return gf


def mark_timeouts(ch):
    for i, r in enumerate(bench):
        if not solved(r):
            dl = ch.series[0].points[i].data_label
            dl.has_text_frame = True
            dl.text_frame.text = f">{r['highs']['time']:.0f} s, not solved"
            dl.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
            dl.text_frame.paragraphs[0].runs[0].font.bold = True
            dl.position = XL_LABEL_POSITION.ABOVE if ch.chart_type == XL_CHART_TYPE.LINE_MARKERS \
                else XL_LABEL_POSITION.OUTSIDE_END


# =========================================================== slide 1: title
s1 = prs.slides[0]
for sh in s1.shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                t = r.text
                t = t.replace("SIH26153", "SIH26119")
                t = t.replace("AI based Network Attack Forecasting from Network Traffic Data",
                              "Indigenous GPU-Accelerated Optimization Solver (Sovereign Alternative to Express / CEPLEX)")
                t = t.replace("Blockchain & Cybersecurity", "Smart Automation")
                r.text = t
        for p in sh.text_frame.paragraphs:
            ppr = p._p.find(qn("a:pPr"))
            if ppr is not None and ppr.get("algn") == "just":
                ppr.set("algn", "l")
                if "Problem Statement Title" in p.text:
                    ppr.find(qn("a:lnSpc")).find(qn("a:spcPct")).set("val", "130000")

# =========================================================== slide 2: idea
s2 = prs.slides[1]
strip(s2)
set_title(s2, "SAMADHAN: a sovereign, GPU-first optimization solver built from scratch", 26)
textbox(s2, 0.45, 1.25, 7.55, 5.6, [
    para("h1", "Proposed Solution (Describe your Idea/Solution/Prototype)"),
    para("h2", "Detailed explanation of the proposed solution"),
    para("b", [("Own solver core, no borrowed engine: ", "MPS reader, scaling and every algorithm "
                "written by us; no CPLEX, Gurobi, Xpress, HiGHS or CBC code inside")]),
    para("b", [("GPU LP engine (working today): ", "restarted primal-dual hybrid gradient (PDLP); needs only "
                "sparse matrix-vector products, so it runs fully on an NVIDIA GPU")]),
    para("b", [("MILP started, QP next: ", f"our branch-and-bound on our own simplex already proves the same "
                f"optimum as HiGHS on {milp_ok}/{len(milp)} refinery MILPs; cuts, heuristics and QP follow")]),
    para("h2", "How it addresses the problem"),
    para("b", [("Built for refinery planning: ", f"solves crude-blending and product-dispatch LPs up to "
                f"{big['n']/1e6:.1f} million variables on a laptop GPU")]),
    para("b", [("Faster where it matters: ", f"on the largest model, {spd_txt(big)} faster than HiGHS with cost "
                f"within {pct(big['acc']['obj_err'])} of optimal; {spd_txt(big, 'fast')} faster in fast mode")]),
    para("h2", "Innovation and uniqueness of the solution"),
    para("b", [("GPU-first, factorisation-free: ", "memory-light and massively parallel, scales past the "
                "limits of simplex and interior point")]),
    para("b", [("Transparent and tunable: ", "open, auditable code that Indian engineers can inspect and "
                "adapt to MRPL's own models")]),
])
time_chart(s2, 8.2, 1.3, 4.85, 3.2, "Solve time, refinery planning LP (log scale)")
caption(s2, 8.25, 4.5, 4.8, 0.45, "Working prototype; same models, same laptop. Each solver credited "
        "with its faster method.")
stat(s2, 8.25, 5.05, 1.5, 1.6, spd_txt(big), f"faster than HiGHS on {big['n']/1e6:.1f}M variables")
stat(s2, 9.9, 5.05, 1.5, 1.6, spd_txt(big, "fast"), "faster in fast mode (1e-4)")
stat(s2, 11.55, 5.05, 1.5, 1.6, pct(worst_acc), "worst cost error vs optimum (1e-6)", size=19)

# =========================================================== slide 3: technical approach
s3 = prs.slides[2]
strip(s3)
textbox(s3, 0.45, 1.2, 12.4, 1.6, [
    para("h1", "Technologies to be used"),
    para("b", [("Solver core: ", "Python 3 + PyTorch CUDA sparse kernels (prototype), moving hot loops to C++/CUDA; "
                "NumPy/SciPy only as data containers")]),
    para("b", [("Interfaces: ", "Python API, command-line tool, MPS/LP files; REST service for plant planning systems")]),
    para("b", [("Hardware: ", "any NVIDIA GPU (tested on RTX 5050 Laptop, 8 GB); the same code runs on CPU-only servers")]),
    para("h1", "Methodology and process for implementation", before=600),
])
# architecture diagram: row of pipeline boxes
bx_y, bx_h, bx_w, gap = 3.0, 1.55, 2.25, 0.27
boxes = [
    ("Model input", "MPS / LP files, Python API, MRPL planning data", True),
    ("Scaling", "Ruiz and Pock-Chambolle equilibration (presolve added in build phase)", True),
    ("GPU LP engine", "restarted PDHG, adaptive steps and restarts, all on CUDA", True),
    ("Polish + crossover", "dual simplex on CPU for exact vertex solutions", False),
    ("Result + audit", "solution, duals, KKT certificate, logs", True),
]
x0 = 0.55
for i, (h, d, done) in enumerate(boxes):
    x = x0 + i * (bx_w + gap)
    card(s3, x, bx_y, bx_w, bx_h, fill=GREEN_BG if done else AMBER_BG, line=GREEN if done else "C98A1B",
         paras=[para("p", [(h, "")], size=13, align="ctr", color=BLUE),
                para("p", d, size=10.5, align="ctr", before=300)], anchor=MSO_ANCHOR.MIDDLE)
    if i < len(boxes) - 1:
        a = s3.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bx_w + 0.03), Inches(bx_y + bx_h / 2 - 0.12),
                                Inches(gap - 0.06), Inches(0.24))
        a.fill.solid(); a.fill.fore_color.rgb = RGBColor.from_string("7F7F7F"); a.line.fill.background()
# MILP / QP layer spanning under engine
lay_y = 4.85
card(s3, 0.55, lay_y, 6.0, 1.45, fill=GREEN_BG, line=GREEN, paras=[
    para("p", [("MILP layer: ", f"branch-and-bound on our own simplex (built; exact on {milp_ok}/{len(milp)} "
                "test MILPs). Next: cuts (Gomory, MIR, cover), presolve, heuristics (feasibility pump, RINS), "
                "GPU LP for large relaxations")], size=11)],
    anchor=MSO_ANCHOR.MIDDLE)
card(s3, 6.75, lay_y, 3.4, 1.45, fill="FFFFFF", line=BLUE, paras=[
    para("p", [("QP layer: ", "primal-dual QP (PDQP) on the same GPU kernels; later MIQP and NLP")], size=11)],
    anchor=MSO_ANCHOR.MIDDLE)
card(s3, 10.35, lay_y, 2.55, 1.45, fill="FFFFFF", line="A6A6A6", paras=[
    para("p", [("Legend", "")], size=11, color=INK),
    para("p", [("■ ", "built and tested today")], size=10.5, color=GREEN, before=300),
    para("p", [("■ ", "planned, SIH build phase")], size=10.5, color="C98A1B", before=200)],
    anchor=MSO_ANCHOR.MIDDLE)
caption(s3, 0.55, 6.42, 12.3, 0.4, "Every stage is our own code. HiGHS is used only as an outside benchmark to check "
        "answers and speed.")

# =========================================================== slide 4: feasibility
s4 = prs.slides[3]
strip(s4)
textbox(s4, 0.45, 1.2, 6.3, 2.1, [
    para("h1", "Analysis of the feasibility of the idea"),
    para("b", [("Already working: ", f"every refinery test LP solved, cost within {pct(worst_acc)} "
                "of the true optimum")], size=12),
    para("b", [("Modest hardware: ", f"{big['n']/1e6:.1f}M variables, {big['nnz']/1e6:.1f}M nonzeros "
                "on one 8 GB laptop GPU")], size=12),
    para("b", [("Proven, published algorithms: ", "PDLP-type methods are peer-reviewed and used in "
                "commercial solvers")], size=12),
    para("b", [("Clear plan: ", "LP done, MILP working on small models; cuts, QP and scale-up next")], size=12),
])
if netlib:
    ok = [r for r in netlib if r.get("status") == "optimal"]
    close = [r for r in ok if r["obj_rel_diff"] <= 1e-3]
    parser_ok = sum(1 for r in netlib if r.get("parser_obj_diff", 1) < 1e-9)
    card(s4, 0.55, 3.3, 6.1, 1.2, paras=[
        para("p", [(f"Netlib: {len(ok)}/{len(netlib)} LPs solved ", f"to 1e-4 relative KKT ({len(close)} within "
                    f"0.1% of optimum); MPS reader matches HiGHS on {parser_ok}/{len(netlib)}. The "
                    f"{len(netlib) - len(ok)} left are ill-conditioned cases for our simplex crossover.")],
             size=11),
        para("p", [(f"MILP: {milp_ok}/{len(milp)} refinery MILPs ", "proven optimal by our branch-and-bound, "
                    "same cost as HiGHS.")], size=11, before=300)], anchor=MSO_ANCHOR.MIDDLE)

# risks table on the right
risks = [
    ("Risk", "How we handle it"),
    ("First-order methods give moderate accuracy", "Restarts and tight KKT checks, then simplex crossover for exact vertices"),
    ("Hard MILPs need years of tuning", "Start with the MIPLIB 'easy' set and refinery MILPs; add cuts and heuristics step by step"),
    ("Degenerate or ill-conditioned models", "Ruiz scaling, float64, presolve; crossover for Netlib's hard "
     "cases (pilot4, perold, greenbea)"),
    ("No GPU at a site", "Same code runs on CPU; GPU only where it is measurably faster"),
]
tbl = s4.shapes.add_table(len(risks), 2, Inches(6.9), Inches(1.3), Inches(6.0), Inches(3.1)).table
tbl.columns[0].width = Inches(2.3); tbl.columns[1].width = Inches(3.7)
for i in range(len(risks)):
    tbl.rows[i].height = Inches(0.36 if i == 0 else 0.62)
for i, (a_, b_) in enumerate(risks):
    for j, txt in enumerate((a_, b_)):
        c = tbl.cell(i, j)
        c.text = txt
        c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(BLUE if i == 0 else ("F7F9FC" if i % 2 else "FFFFFF"))
        c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.04)
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10.5); r.font.name = "Arial"; r.font.bold = (i == 0 or j == 0)
                r.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else INK)
# benchmark table across the bottom
hdr = ["Refinery LP", "Variables", "Nonzeros", "HiGHS (best)", "Accurate 1e-6", "Cost error",
       "Speed-up", "Fast 1e-4", "Cost error", "Speed-up"]
rows = [hdr] + [[f"Size {r['size']}", f"{r['n']:,}", f"{r['nnz']:,}", htime_txt(r),
                 fmt_s(r['acc']['time']), pct(r['acc']['obj_err']), spd_txt(r),
                 fmt_s(r['fast']['time']), pct(r['fast']['obj_err']), spd_txt(r, 'fast')] for r in bench]
tb = s4.shapes.add_table(len(rows), len(hdr), Inches(0.45), Inches(4.65), Inches(12.45), Inches(0.3 * len(rows))).table
for j, wd in enumerate((1.3, 1.3, 1.3, 1.35, 1.35, 1.2, 1.05, 1.2, 1.2, 1.2)):
    tb.columns[j].width = Inches(wd)
for i, row in enumerate(rows):
    for j, txt in enumerate(row):
        c = tb.cell(i, j)
        c.text = txt
        c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string(BLUE if i == 0 else ("F7F9FC" if i % 2 else "FFFFFF"))
        c.margin_top = c.margin_bottom = Inches(0.03)
        for p in c.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            for r in p.runs:
                r.font.size = Pt(10); r.font.name = "Arial"; r.font.bold = (i == 0 or j in (6, 9))
                r.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else INK)
caption(s4, 0.45, 4.7 + 0.3 * len(rows) + 0.05, 12.4, 0.4,
        "Measured today on one laptop (RTX 5050 8 GB GPU). HiGHS: faster of dual simplex and interior point, all "
        "CPU threads. SAMADHAN: faster of its two GPU methods, at 1e-6 or 1e-4 relative KKT. Cost error vs HiGHS optimum.")

# =========================================================== slide 5: impact
s5 = prs.slides[4]
strip(s5)
textbox(s5, 0.45, 1.2, 6.5, 2.4, [
    para("h1", "Potential impact on the target audience"),
    para("b", [("MRPL and other refineries: ", "faster crude selection, blending and dispatch plans; "
                "more what-if scenarios per planning cycle")]),
    para("b", [("Power, logistics, steel, railways: ", "the same engine serves unit commitment, routing "
                "and production planning")]),
    para("b", [("Strategic users: ", "runs air-gapped and on-premise, with no licence server and no data "
                "leaving the site")]),
])
textbox(s5, 0.45, 3.45, 6.8, 0.4, [para("h1", "Benefits of the solution (social, economic, environmental, etc.)", size=14.5)])
cards = [
    ("Economic", "No per-core or per-user licence fees; GPU speed-ups cut planning time"),
    ("Sovereignty", "Indian-owned solver code that can be audited, changed and certified"),
    ("Environmental", "Better crude and energy plans mean less fuel, flaring and waste"),
    ("Social", "Builds Indian skills in optimisation software, a rare capability"),
]
for i, (h, d) in enumerate(cards):
    x = 0.55 + (i % 2) * 3.27
    y = 4.05 + (i // 2) * 1.42
    card(s5, x, y, 3.13, 1.3, paras=[para("p", [(h, "")], size=14, color=BLUE),
                                     para("p", d, size=11.5, before=300)])
# scaling line chart on the right
cd = CategoryChartData()
cd.categories = [label(r) for r in bench]
cd.add_series("HiGHS (best)", [round(r["highs"]["time"], 2) for r in bench])
cd.add_series("SAMADHAN accurate", [round(r["acc"]["time"], 2) for r in bench])
cd.add_series("SAMADHAN fast", [round(r["fast"]["time"], 2) for r in bench])
gf = s5.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(7.3), Inches(1.3), Inches(5.7), Inches(3.6), cd)
ch = gf.chart
ch.has_title = True; ch.chart_title.text_frame.text = "Solve time vs model size (seconds, log scale)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
ch.legend.font.size = Pt(10)
for s, col in zip(ch.series, SERIES_COLORS):
    s.format.line.color.rgb = RGBColor.from_string(col); s.format.line.width = Pt(2.5)
    s.marker.format.fill.solid(); s.marker.format.fill.fore_color.rgb = RGBColor.from_string(col)
    s.marker.format.line.color.rgb = RGBColor.from_string(col); s.smooth = False
ch.value_axis.tick_labels.font.size = Pt(9)
ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string("E0E0E0")
ch.category_axis.tick_labels.font.size = Pt(10)
ch.category_axis.has_title = True; ch.category_axis.axis_title.text_frame.text = "Model size (variables)"
ch.category_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
log_axis(ch)
mark_timeouts(ch)
caption(s5, 7.35, 4.95, 5.6, 0.6, f"Refinery planning LPs from {label(bench[0])} to {label(big)} variables. "
        "The gap widens with size, which is where real plant-wide models live.")
card(s5, 7.35, 5.55, 5.6, 1.1, fill=CARD, paras=[
    para("p", [("Every licence we replace ", "is money saved each year, and every model we solve in-house "
                "stays under Indian control.")], size=12)], anchor=MSO_ANCHOR.MIDDLE)

# =========================================================== slide 6: references
s6 = prs.slides[5]
strip(s6)
ref = lambda t, d: para("b", [(t, d)], size=11.5, before=250)
import qrcode
qr_path = HERE / "qr_repo.png"
qrcode.make(REPO, border=1).save(qr_path)
card(s6, 9.55, 1.35, 3.35, 3.2, fill=CARD)
s6.shapes.add_picture(str(qr_path), Inches(10.3), Inches(1.5), Inches(1.85), Inches(1.85))
tb = textbox(s6, 9.6, 3.45, 3.25, 1.05, [
    para("p", [("Source code and benchmarks", "")], size=12, align="ctr", color=BLUE),
    para("p", "github.com/NAMAN-THAKUR-1944/", size=9.5, align="ctr", before=150),
    para("p", "SIH2026-SIH26119-SAMADHAN", size=9.5, align="ctr")])
card(s6, 9.55, 4.75, 3.35, 1.95, fill="FFFFFF", line="A6A6A6", paras=[
    para("p", [("Reproduce every number", "")], size=12, color=BLUE),
    para("p", "python -m samadhan demo --size L", size=9.5, before=300),
    para("p", "python bench_final.py S M L XL", size=9.5, before=150),
    para("p", "python bench_netlib.py", size=9.5, before=150),
    para("p", "python bench_milp.py", size=9.5, before=150)], anchor=MSO_ANCHOR.MIDDLE)
textbox(s6, 0.45, 1.2, 8.9, 5.6, [
    para("h1", "Details / Links of the reference and research work"),
    para("h2", "LP on GPUs (our core)"),
    ref("PDLP: ", "Applegate, Díaz, Hinder, Lu, Lubin, O'Donoghue, Schudy, Practical Large-Scale Linear Programming "
        "using Primal-Dual Hybrid Gradient, NeurIPS 2021"),
    ref("cuPDLP: ", "Lu & Yang, cuPDLP.jl: A GPU Implementation of Restarted PDHG for Linear Programming, arXiv:2311.12180 (2023)"),
    ref("PDHG and scaling: ", "Chambolle & Pock, J. Math. Imaging Vision 40 (2011); Ruiz, A scaling algorithm to "
        "equilibrate both rows and columns norms in matrices, RAL-TR-2001-034"),
    para("h2", "MILP and QP (build phase)"),
    ref("Branch-and-cut: ", "Achterberg, Constraint Integer Programming, PhD thesis, TU Berlin (2007); "
        "Bixby, A Brief History of Linear and Mixed-Integer Programming Computation, Doc. Math. (2012)"),
    ref("Dual simplex: ", "Huangfu & Hall, Parallelizing the dual revised simplex method, Math. Prog. Comp. 10 (2018)"),
    ref("QP: ", "Lu & Yang, A Practical and Optimal First-Order Method for Large-Scale Convex QP (PDQP), arXiv:2311.07710; "
        "Stellato et al., OSQP, Math. Prog. Comp. 12 (2020)"),
    para("h2", "Benchmarks the PS asks for"),
    ref("MIPLIB 2017: ", "Gleixner et al., Math. Prog. Comp. 13 (2021), miplib.zib.de"),
    ref("Netlib LP and Mittelmann benchmarks: ", "netlib.org/lp  •  plato.asu.edu/bench.html"),
])

prs.save(str(HERE / "SAMADHAN_SIH26119_Idea.pptx"))
print("saved")
